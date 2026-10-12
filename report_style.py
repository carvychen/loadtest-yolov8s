"""Word 报告共用版式 (make_report.py 一期 / make_report_fp8.py 二期)。风格: 专业、精炼、清晰。"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LATIN = "Calibri"
CJK = "Microsoft YaHei"
INK = RGBColor(0x21, 0x21, 0x21)
ACCENT = RGBColor(0x12, 0x57, 0xA0)
MUTED = RGBColor(0x60, 0x6A, 0x76)


def set_cjk(run):
    run.font.name = LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)


def base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    for i, sz in ((1, 16), (2, 12.5)):
        h = doc.styles[f"Heading {i}"]
        h.font.name = LATIN
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = ACCENT if i == 1 else INK
        h.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
        h.paragraph_format.space_before = Pt(14 if i == 1 else 10)
        h.paragraph_format.space_after = Pt(4)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def cell_text(cell, text, bold=False, color=None, size=9.5, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(text)
    set_cjk(r)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = color


def make_table(doc, headers, rows, widths=None, highlight_row=None):
    """highlight_row 可为单个下标或下标集合。"""
    hl_rows = set() if highlight_row is None else (
        {highlight_row} if isinstance(highlight_row, int) else set(highlight_row))
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9.5,
                  align="left" if j == 0 else "center")
        shade(c, "1257A0")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        hl = i in hl_rows
        for j, val in enumerate(row):
            cell_text(cells[j], val, bold=hl, size=9.5,
                      align="left" if j == 0 else "center")
            if hl:
                shade(cells[j], "E7F0FA")
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    return t


def para(doc, text, size=10.5, color=None, italic=False, bold=False, after=6, align="left"):
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_cjk(r)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.bold = bold
    if color:
        r.font.color.rgb = color
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_cjk(r)
    r.font.size = Pt(size)


def new_doc(title, subtitle, basis):
    """建文档 + 版心边距 + 标题/副标题/测试基础三行。"""
    doc = Document()
    base_style(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(0.9)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run(title)
    set_cjk(r)
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    para(doc, subtitle, size=11, color=MUTED, after=2)
    para(doc, basis, size=9, color=MUTED, after=10)
    return doc
